from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw

from app.financial_helpers import format_money
from app.pdf_rendering import load_font, wrap_text


SOURCE_TITLES = {
    "debtors": "Должники",
    "csi": "ЧСИ",
    "incoming_claims": "Входящая корреспонденция: претензии",
    "incoming_other": "Входящая корреспонденция: прочее",
}
COUNTRY_NAMES = {"kz": "Казахстан", "uz": "Узбекистан"}


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    kind: str = "text"
    width: int = 200


@dataclass
class Report:
    title: str
    country: str
    columns: list[Column]
    rows: list[dict]
    period: str | None = None

    def totals(self) -> dict[str, Decimal]:
        return {
            col.key: sum((Decimal(str(row.get(col.key) or 0)) for row in self.rows), Decimal("0")).quantize(Decimal("0.01"))
            for col in self.columns if col.kind == "money"
        }


def debtor_columns(country: str) -> list[Column]:
    currency = "тг" if country == "kz" else "сум"
    return [
        Column("entry_date_iso", "Дата внесения", "date", 140),
        Column("contract_date", "Дата контракта", "date", 140),
        Column("category", "Категория", width=240),
        Column("client_name", "ФИО клиента", width=300),
        Column("contract_number", "Номер договора", width=220),
        Column("last_missed_payment_date_iso", "Дата первого неисполненного платежа", "date", 210),
        Column("company", "Компания", width=240),
        Column("city", "Город", width=160),
        Column("court", "Суд", width=300),
        Column("claim_sent", "Претензия", "bool", 130),
        Column("claim_sent_date_iso", "Дата отправки претензии", "date", 160),
        Column("claim_sent_days", "Дней с отправки претензии", "number", 150),
        Column("debt_days", "Дней долга", "number", 130),
        Column("debt_amount", f"Сумма долга ({currency})", "money", 170),
        Column("penalty_amount", f"Пеня ({currency})", "money", 170),
        Column("state_duty_amount", f"Госпошлина ({currency})", "money", 170),
        Column("total_amount", f"Общая сумма ({currency})", "money", 180),
        Column("lawsuit_sent", "Направлен иск", "bool", 140),
        Column("lawsuit_sent_date_iso", "Дата отправки иска", "date", 150),
        Column("lawsuit_accepted", "Иск принят", "bool", 130),
        Column("hearing_date_iso", "Дата заседания", "date", 150),
        Column("decision_exists", "Есть решение", "bool", 130),
        Column("decision", "Решение", width=200),
        Column("decision_payout", f"Сумма выплаты по решению ({currency})", "money", 190),
        Column("received_amount", f"Получено ({currency})", "money", 170),
        Column("comment", "Комментарий", width=330),
        Column("case_number", "Номер дела", width=190),
        Column("case_court", "Суд по делу", width=300),
    ]


def incoming_columns(claims: bool) -> list[Column]:
    columns = [
        Column("received_date_iso", "Дата получения письма", "date", 160),
        Column("receive_method", "Способ получения", width=180),
        Column("company", "Компания", width=240),
        Column("client_name", "ФИО клиента", width=300),
        Column("authority_display", "Суд или другой орган", width=300),
        Column("contract_number", "Номер договора", width=220),
        Column("category", "Категория", width=250),
    ]
    if claims:
        columns.extend([
            Column("responsible_person", "Ответственный за подготовку ответа", width=240),
            Column("response_text", "Ответ", width=220),
            Column("response_date_iso", "Дата ответа", "date", 150),
            Column("comment", "Комментарий", width=330),
        ])
    return columns


def table_report(source: str, country: str, rows: list[dict]) -> Report:
    columns = incoming_columns(source == "incoming_claims") if source.startswith("incoming_") else debtor_columns(country)
    prepared = [dict(row) for row in rows]
    if source in ("debtors", "csi"):
        by_id = {row["id"]: row for row in prepared}
        for row in prepared:
            parent = by_id.get(row.get("parent_debtor_id"))
            if parent is not None:
                # The UI shares these cells with the parent; do not count them twice.
                for key in ("client_name", "contract_number", "company", "city", "court", "case_court", "entry_date_iso", "contract_date", "last_missed_payment_date_iso"):
                    row[key] = parent.get(key)
                for key in ("debt_amount", "penalty_amount", "state_duty_amount", "total_amount", "debt_days"):
                    row[key] = None
    return Report(SOURCE_TITLES[source], country, columns, prepared)


def payment_report(source: str, country: str, rows: list[dict], date_from: date, date_to: date) -> Report:
    currency = "тг" if country == "kz" else "сум"
    return Report(
        f"Рапорт по выплатам: {SOURCE_TITLES[source]}", country,
        [Column("payment_date", "Дата оплаты", "date", 180), Column("contract_number", "Номер договора клиента", width=270),
         Column("client_name", "ФИО клиента", width=380), Column("amount", f"Сумма ({currency})", "money", 210)],
        rows, f"{date_from:%d.%m.%Y} - {date_to:%d.%m.%Y}",
    )


def cell_value(column: Column, value):
    if value is None or value == "":
        return None
    if column.kind == "date":
        if isinstance(value, date):
            return value
        for pattern in ("%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(str(value), pattern).date()
            except ValueError:
                pass
        return str(value)
    if column.kind == "bool":
        return "Да" if value else "Нет"
    return value


def output_path(directory: Path, extension: str) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    return directory / f"export_{uuid4().hex}.{extension}"


def render_xlsx(report: Report, directory: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Рапорт"
    sheet.sheet_view.showGridLines = False
    sheet.append([report.title])
    sheet.append([f"Страна: {COUNTRY_NAMES[report.country]}"])
    if report.period:
        sheet.append([f"Период: {report.period}"])
    sheet.append([f"{'Поступлений' if report.period else 'Записей'}: {len(report.rows)}"])
    totals = report.totals()
    if report.period:
        sheet.append(["Общая сумма", float(totals["amount"])])
        sheet.cell(sheet.max_row, 2).number_format = '#,##0.00'
    header_row = sheet.max_row + 2
    sheet.append([])
    sheet.append([column.label for column in report.columns])
    for record in report.rows:
        sheet.append([cell_value(column, record.get(column.key)) for column in report.columns])
    last_data_row = sheet.max_row
    if totals:
        sheet.append(["Итого" if index == 0 else (float(totals[column.key]) if column.key in totals else None)
                      for index, column in enumerate(report.columns)])

    for row in sheet:
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)
                cell.data_type = "s"
    sheet["A1"].font = Font(name="Arial", size=14, bold=True)
    for row_number in range(1, header_row):
        sheet.cell(row_number, 1).alignment = Alignment(vertical="center", wrap_text=False)
        sheet.row_dimensions[row_number].height = 21
    sheet.row_dimensions[1].height = 24
    for cell in sheet[header_row]:
        cell.fill = PatternFill("solid", fgColor="18334A")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 44
    for index, column in enumerate(report.columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(18, column.width / 7)
        for row_number in range(header_row + 1, sheet.max_row + 1):
            cell = sheet.cell(row_number, index)
            if column.kind == "money":
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            elif column.kind == "date":
                cell.number_format = "dd.mm.yyyy"
            elif column.kind == "number":
                cell.number_format = '#,##0'
    for row_number in range(header_row + 1, last_data_row + 1):
        line_count = max((sum(max(1, (len(line) + int(sheet.column_dimensions[cell.column_letter].width) - 1) // int(sheet.column_dimensions[cell.column_letter].width))
                              for line in str(cell.value or "").split("\n")) for cell in sheet[row_number]), default=1)
        sheet.row_dimensions[row_number].height = min(409, max(24, line_count * 15))
        if (row_number - header_row) % 2 == 0:
            for cell in sheet[row_number]:
                cell.fill = PatternFill("solid", fgColor="F0F4F7")
    if totals:
        for cell in sheet[sheet.max_row]:
            cell.font = Font(name="Arial", size=10, bold=True)
            cell.border = Border(top=Side(style="thin", color="18334A"))
        sheet.row_dimensions[sheet.max_row].height = 26
    sheet.freeze_panes = f"{'C' if report.period else 'A'}{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(report.columns))}{last_data_row}"
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    path = output_path(directory, "xlsx")
    workbook.save(path)
    workbook.close()
    return path


def render_pdf(report: Report, directory: Path) -> Path:
    # Keep every column on one landscape A4 page; extra pixels preserve detail when zooming.
    scale = max(1.0, sum(column.width for column in report.columns) / (1754 - 120))
    page_width, page_height = round(1754 * scale), round(1240 * scale)
    margin, bottom = round(60 * scale), page_height - round(70 * scale)
    regular = [Path(r"C:\Windows\Fonts\arial.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")]
    bold = [Path(r"C:\Windows\Fonts\arialbd.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")]
    font = load_font(18, regular_candidates=regular, bold_candidates=bold)
    header_font = load_font(18, bold=True, regular_candidates=regular, bold_candidates=bold)
    meta_font = load_font(round(18 * scale), regular_candidates=regular, bold_candidates=bold)
    title_font = load_font(round(27 * scale), bold=True, regular_candidates=regular, bold_candidates=bold)
    columns = report.columns
    anchors = [column for column in columns if column.key in ("client_name", "contract_number")]
    factor = (page_width - margin * 2) / sum(column.width for column in columns)
    widths = [round(column.width * factor) for column in columns]
    widths[-1] += page_width - margin * 2 - sum(widths)
    totals = report.totals()
    path = output_path(directory, "pdf")
    page_number = 0

    def finish_page(image):
        nonlocal page_number
        page_number += 1
        ImageDraw.Draw(image).text(
            (margin, page_height - round(42 * scale)), f"Страница {page_number}",
            fill="#666666", font=meta_font,
        )
        # Append one page at a time to keep large, unpaginated exports bounded in memory.
        image.save(path, "PDF", resolution=150 * scale, append=page_number > 1)
        image.close()

    def text(column, value):
        value = cell_value(column, value)
        if value is None:
            return ""
        if isinstance(value, date):
            return f"{value:%d.%m.%Y}"
        return format_money(float(value)) if column.kind == "money" else str(value)

    def new_page():
        image = Image.new("RGB", (page_width, page_height), "white")
        draw = ImageDraw.Draw(image)
        draw.text((margin, round(48 * scale)), report.title, fill="#18334a", font=title_font)
        meta = f"Страна: {COUNTRY_NAMES[report.country]}    {'Поступлений' if report.period else 'Записей'}: {len(report.rows)}"
        if report.period:
            meta += f"    Период: {report.period}    Общая сумма: {format_money(float(totals['amount']))}"
        draw.text((margin, round(90 * scale)), meta, fill="#222222", font=meta_font)
        header_lines = [wrap_text(draw, column.label, header_font, width - 16) for column, width in zip(columns, widths)]
        y, x = round(132 * scale), margin
        height = max(map(len, header_lines)) * 25 + 20
        for width, lines in zip(widths, header_lines):
            draw.rectangle((x, y, x + width, y + height), fill="#18334a", outline="white")
            for index, line in enumerate(lines):
                draw.text((x + 8, y + 10 + index * 25), line, fill="white", font=header_font)
            x += width
        return image, draw, y + height

    image, draw, y = new_page()
    for record_index, record in enumerate(report.rows):
        lines = [sum((wrap_text(draw, paragraph, font, width - 16) for paragraph in text(column, record.get(column.key)).splitlines()), []) or [""]
                 for column, width in zip(columns, widths)]
        offset, line_count = 0, max(map(len, lines))
        while offset < line_count:
            available_lines = (bottom - y - 16) // 25
            if available_lines < 1:
                finish_page(image)
                image, draw, y = new_page()
                continue
            take = min(available_lines, line_count - offset)
            height, x = take * 25 + 16, margin
            for column, width, cell_lines in zip(columns, widths, lines):
                draw.rectangle((x, y, x + width, y + height), fill="#f0f4f7" if record_index % 2 else "white", outline="#c9d3dc")
                visible_lines = cell_lines[:take] if offset and column in anchors else cell_lines[offset:offset + take]
                for line_index, line in enumerate(visible_lines):
                    line_x = x + width - 8 - draw.textlength(line, font=font) if column.kind == "money" else x + 8
                    draw.text((line_x, y + 8 + line_index * 25), line, fill="#111111", font=font)
                x += width
            y += height
            offset += take
    if not report.rows:
        draw.text((margin + 8, y + 18), "Нет данных по выбранным условиям.", fill="#333333", font=font)
        y += 60
    for column in columns:
        if column.key in totals:
            label = f"Итого - {column.label}: {format_money(float(totals[column.key]))}"
            if y + 40 > bottom:
                finish_page(image)
                image, draw, y = new_page()
            draw.text((margin, y + 12), label, fill="#18334a", font=header_font)
            y += 38
    finish_page(image)
    return path
